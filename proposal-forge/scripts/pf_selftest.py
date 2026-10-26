#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
proposal-forge / pf_selftest.py
===============================
安装自检：造一份临时文档，把全链路跑一遍，逐项断言。

    python scripts/pf_selftest.py [--keep]

不读任何外部素材，全部现场生成，适合下载后第一时间验证环境。
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import pf_docx as D
import pf_render as R
import pf_style as S

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  [ok]   " if cond else "  [FAIL] ") + name + (("  → " + detail) if detail and not cond else ""))


def make_docx(path):
    """造一份结构类似申报书的文档：封面段 + 标题 + 正文 + 表格 + 图 + 题注乱序。"""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("项目名称  测试用申报书")
    doc.add_paragraph("项目负责人  张三")

    # 正文塞进一个 1×1 表格 —— 复刻中文申报书的常见结构
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]

    cell.paragraphs[0].text = "1 研究背景"
    p = cell.add_paragraph("本项目围绕测试场景展开，样本量为 103 份有效问卷。")

    img = R.chart_grouped_bar(
        str(Path(path).with_name("t_bar.png")), ["甲", "乙", "丙"],
        [("指标一", [72.1, 58.4, 40.2]), ("指标二", [61.3, 44.7, 39.8])],
        ylabel="得分", figsize=(4.2, 2.6))
    cell.add_paragraph().add_run().add_picture(img, width=Cm(12))

    # 故意把题注写成乱序 + 带空格，验证重排能修
    cell.add_paragraph("图 0 测试图一")
    cell.add_paragraph("2 研究方法")
    cell.add_paragraph("采用问卷调查与实地访谈相结合的方式。")
    cell.add_paragraph("图7 测试图二")
    cell.add_paragraph("参考文献")
    cell.add_paragraph("[1] 张三. 测试文献[J]. 测试期刊, 2023, 12(3): 1-9.")
    doc.save(path)
    return path


def main():
    keep = "--keep" in sys.argv
    tmp = Path(tempfile.mkdtemp(prefix="pf_selftest_"))
    print("工作目录:", tmp)

    print("\n[1/7] 生成测试文档")
    src = make_docx(str(tmp / "src.docx"))
    check("造出 docx", os.path.exists(src))

    print("[2/7] 结构探测")
    doc = D.open_doc(str(src))
    st = D.text_stats(doc)
    check("穿透表格统计到正文", st["paragraphs"] >= 5, "paragraphs=%d" % st["paragraphs"])
    check("统计到内联图片", st["images_inline"] == 1, "images=%d" % st["images_inline"])

    print("[3/7] 图表编号重排")
    fig, tab, changed = D.renumber_captions(doc)
    check("识别出 2 张图", fig == 2, "fig=%d" % fig)
    cap = D.audit_captions(doc)
    check("编号重排后连续", cap["fig_issues"][0] == [] and cap["fig_issues"][1] == [],
          str(cap["fig_issues"]))

    print("[4/7] 表格防裁 + 图片宽度")
    n_split = D.fix_table_split(doc)
    check("表格行加 cantSplit", n_split >= 1, "n=%d" % n_split)
    out = D.save_as(doc, src, str(tmp / "work.docx"))
    check("另存成功", os.path.exists(out))
    d2 = D.open_doc(out)
    widths = D.image_width_report(d2)
    check("读到图片显示宽度", len(widths) == 1 and widths[0]["cm"] > 5,
          str(widths)[:80])

    print("[5/7] 交付自检")
    rep = D.qa_report(out)
    md = D.format_qa_markdown(rep)
    check("自检报告生成", len(md) > 300, "len=%d" % len(md))
    check("文献编号对应", rep["refs"]["missing_entries"] == [],
          str(rep["refs"]["missing_entries"]))

    print("[6/7] 去 AI 味检测")
    ai_text = ("随着乡村振兴战略的不断深入推进，本项目具有深远影响。"
               "业内普遍认为，相关研究表明该模式具有很强的可行性。"
               "值得注意的是，我们进行了深入的研究，通过调研的方式掌握了情况。"
               "不仅提升了认知，而且增强了能力，同时促进了协同——这一点至关重要。")
    a = S.analyze(ai_text)
    clean = "样本 103 份，有效 100 份。认知分 75.3，执行分 40.2，落差 35.1 分。"
    b = S.analyze(clean)
    check("AI 文本得分低于干净文本", a["score"] < b["score"],
          "ai=%d clean=%d" % (a["score"], b["score"]))
    check("命中铺垫/拔高/归因三类", len({f["code"] for f in a["findings"]}) >= 3,
          str([f["code"] for f in a["findings"]]))

    print("[7/7] 图表与图像")
    d = tmp / "fig"
    d.mkdir()
    made = [
        R.chart_grouped_bar(str(d / "bar.png"), ["甲", "乙"], [("A", [1, 2]), ("B", [2, 1])]),
        R.chart_dumbbell(str(d / "dumbbell.png"), ["甲", "乙"], [10, 20], [30, 25]),
        R.chart_gantt(str(d / "gantt.png"), [("任务", 1, 6, 6, "交付")], 8),
        R.chart_donut(str(d / "donut.png"), ["甲", "乙", "丙"], [47, 18, 17], center_text="82"),
        R.chart_heatmap(str(d / "heatmap.png"), ["甲", "乙"], ["X", "Y"], [[1, 2], [3, 4]]),
        R.chart_slope(str(d / "slope.png"), {"甲": 1, "乙": 2}, {"甲": 3, "乙": 4}),
        R.chart_framework(str(d / "framework.png"),
                          [("层一", ["块一", "块二"]), ("层二", ["块三"])]),
    ]
    check("七种图型全部出图", all(os.path.exists(x) and os.path.getsize(x) > 3000 for x in made),
          str([os.path.basename(x) for x in made if os.path.getsize(x) < 3000]))

    from PIL import Image
    rgba = Image.new("RGBA", (400, 200), (0, 0, 0, 0))
    rgba.paste((30, 90, 120, 255), (50, 50, 350, 150))
    p_rgba = tmp / "rgba.png"
    rgba.save(p_rgba)
    p_white = R.white_bg(str(p_rgba), str(tmp / "white.png"))
    check("透明图压白底成 RGB", Image.open(p_white).mode == "RGB",
          Image.open(p_white).mode)

    p_stack = R.hstack([p_white, p_white], str(tmp / "stack.png"))
    im_stack = Image.open(p_stack)
    check("横排拼接宽度翻倍", im_stack.width > Image.open(p_white).width,
          "%d vs %d" % (im_stack.width, Image.open(p_white).width))

    hues = R.dominant_hues(p_white)
    check("能识别主色相", len(hues) >= 1, str(hues))
    p_re = R.recolor_family(p_white, str(tmp / "recolor.png"), [200.0], [20.0])
    check("配色迁移出图", os.path.exists(p_re))

    print("\n" + "=" * 52)
    print("通过 %d 项，失败 %d 项" % (len(PASS), len(FAIL)))
    if FAIL:
        print("失败项：")
        for f in FAIL:
            print("  -", f)
    print("工作目录:", tmp if keep else "（已清理）")
    if not keep:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
